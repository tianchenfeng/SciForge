from __future__ import annotations

from pathlib import Path
import sys

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from modules.module_gas_fit.service import GasFitService


def main():
    base = Path(__file__).resolve().parent
    input_file = base / "demo_gas_fit.xlsx"
    output_dir = base / "output_gas_fit_stability"
    df = pd.DataFrame({
        "Sample": ["C1", "C2", "T1", "BAD_EMPTY", "BAD_TEXT"],
        "0": [0, 0, 0, None, "x"],
        "2": [5, 6, 4, None, "bad"],
        "4": [12, 13, 9, None, "bad"],
        "note": ["ok", "ok", "ok", "skip", "skip"],
        "8": [25, 26, 20, None, "bad"],
        "12": [38, 39, 32, None, "bad"],
        "24": [60, 61, 55, None, "bad"],
        "36": [70, 71, 66, None, "bad"],
        "48": [75, 76, 72, None, "bad"],
    })
    df.to_excel(input_file, index=False)
    result = GasFitService.run(str(input_file), str(output_dir), "lag", logger=print)

    assert Path(result["excel"]).exists()
    assert Path(result["summary_plot"]).exists()
    assert result["success_count"] == 3
    assert result["failed_count"] == 2
    assert len(result["plots"]) == 3

    workbook = load_workbook(result["excel"], read_only=True)
    assert "fit_results" in workbook.sheetnames
    assert "failed_samples" in workbook.sheetnames
    failed_sheet = workbook["failed_samples"]
    assert failed_sheet.max_row == 3

    print(f"success_count={result['success_count']}")
    print(f"failed_count={result['failed_count']}")
    print(f"excel_path={result['excel']}")
    print(f"plot_dir={output_dir / 'plots'}")
    print("module_gas_fit test passed")


if __name__ == "__main__":
    main()
