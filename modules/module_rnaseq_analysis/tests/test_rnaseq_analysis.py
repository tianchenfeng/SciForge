from __future__ import annotations

from pathlib import Path
import sys

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from modules.module_rnaseq_analysis.service import RNASeqAnalysisService


def main():
    base = Path(__file__).resolve().parent
    output_dir = base / "output_rnaseq_analysis"
    deg_file = base / "demo_rnaseq_analysis_deg.xlsx"
    annotation_file = base / "demo_rnaseq_analysis_annotation.xlsx"
    go_file = base / "demo_rnaseq_analysis_go.xlsx"
    kegg_file = base / "demo_rnaseq_analysis_kegg.xlsx"
    expression_file = base / "demo_rnaseq_analysis_expression.xlsx"

    pd.DataFrame({
        "Gene": ["GeneA", "GeneB", "GeneC", "GeneD", "GeneE", "GeneF"],
        "log2FoldChange": [2.4, -1.8, 0.4, 1.2, -2.2, 1.6],
        "PValue": [0.001, 0.002, 0.5, 0.01, 0.0005, 0.004],
        "FDR": [0.01, 0.03, 0.7, 0.04, 0.02, 0.03],
        "baseMean": [100, 80, 20, 50, 70, 120],
    }).to_excel(deg_file, index=False)

    pd.DataFrame({
        "Gene": ["GeneA", "GeneB", "GeneD", "GeneE", "GeneF"],
        "description": ["cellulose synthase", "unknown", "starch pathway", "fermentation related", "cell wall regulator"],
        "pathway": ["cell wall", "stress", "starch", "fermentation", "cellulose"],
        "GO": ["GO:0009832", "GO:0006950", "GO:0005982", "GO:0008152", "GO:0071555"],
        "KEGG": ["map00500", "map04075", "map00500", "map00680", "map00520"],
    }).to_excel(annotation_file, index=False)

    pd.DataFrame({
        "ID": ["GO:0009832", "GO:0071555", "GO:0005982"],
        "Description": ["plant-type cell wall biogenesis", "cell wall organization", "starch metabolic process"],
        "pvalue": [0.0001, 0.002, 0.01],
        "padj": [0.005, 0.03, 0.04],
        "geneID": ["GeneA/GeneF", "GeneA/GeneF", "GeneD"],
    }).to_excel(go_file, index=False)

    pd.DataFrame({
        "ID": ["map00500", "map00680", "map00520"],
        "Description": ["Starch and sucrose metabolism", "Methane metabolism", "Amino sugar metabolism"],
        "pvalue": [0.0002, 0.006, 0.02],
        "padj": [0.006, 0.04, 0.08],
        "geneID": ["GeneA/GeneD", "GeneE", "GeneF"],
    }).to_excel(kegg_file, index=False)

    pd.DataFrame({
        "Gene": ["GeneA", "GeneB", "GeneD", "GeneE", "GeneF"],
        "C1": [10, 80, 20, 90, 15],
        "C2": [12, 75, 22, 88, 14],
        "T1": [45, 30, 60, 20, 50],
        "T2": [48, 28, 63, 18, 55],
    }).to_excel(expression_file, index=False)

    result = RNASeqAnalysisService.run(
        str(deg_file),
        str(output_dir),
        log2fc_threshold=1.0,
        padj_threshold=0.05,
        keywords="cellulose, starch, fermentation",
        annotation_file=str(annotation_file),
        go_file=str(go_file),
        kegg_file=str(kegg_file),
        expression_file=str(expression_file),
        logger=print,
    )

    required_sheets = [
        "all_deg",
        "filtered_deg",
        "up_genes",
        "down_genes",
        "deg_with_annotation",
        "filtered_deg_with_annotation",
        "up_genes_with_annotation",
        "down_genes_with_annotation",
        "candidate_genes_ranked",
        "top_GO",
        "significant_GO",
        "keyword_GO",
        "top_KEGG",
        "significant_KEGG",
        "keyword_KEGG",
        "pathway_gene_map",
    ]
    assert Path(result["excel"]).exists()
    assert Path(result["word"]).exists()
    assert Path(result["volcano"]).exists()
    assert Path(result["candidate_heatmap"]).exists()
    assert Path(result["top_go_barplot"]).exists()
    assert Path(result["top_kegg_barplot"]).exists()

    workbook = load_workbook(result["excel"], read_only=True)
    for sheet in required_sheets:
        assert sheet in workbook.sheetnames

    candidates = pd.read_excel(result["excel"], sheet_name="candidate_genes_ranked")
    pathway_map = pd.read_excel(result["excel"], sheet_name="pathway_gene_map")
    assert not candidates.empty
    assert not pathway_map.empty
    assert result["up_count"] == 3
    assert result["down_count"] == 2

    print(f"excel_path={result['excel']}")
    print(f"word_path={result['word']}")
    print(f"volcano_path={result['volcano']}")
    print(f"heatmap_path={result['candidate_heatmap']}")
    print(f"go_barplot_path={result['top_go_barplot']}")
    print(f"kegg_barplot_path={result['top_kegg_barplot']}")
    print("module_rnaseq_analysis test passed")


if __name__ == "__main__":
    main()
