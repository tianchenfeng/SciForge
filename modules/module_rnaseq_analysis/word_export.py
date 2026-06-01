from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document


def export_word_report(summary: dict, candidate_df: pd.DataFrame, top_go: pd.DataFrame, top_kegg: pd.DataFrame, output_path: str | Path) -> str:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading("转录组结果综合分析报告", level=1)
    doc.add_paragraph(f"差异基因总数：{summary.get('filtered_count', 0)}")
    doc.add_paragraph(f"上调基因数：{summary.get('up_count', 0)}")
    doc.add_paragraph(f"下调基因数：{summary.get('down_count', 0)}")

    doc.add_heading("候选基因排序", level=2)
    _add_table(doc, candidate_df.head(20), ["Gene", "log2FC", "padj", "KeywordHits", "PathwayHits", "CandidateScore"])

    doc.add_heading("Top GO 富集结果", level=2)
    _add_table(doc, top_go.head(10), ["ID", "Term", "padj", "GeneCount", "KeywordHits"])

    doc.add_heading("Top KEGG 富集结果", level=2)
    _add_table(doc, top_kegg.head(10), ["ID", "Term", "padj", "GeneCount", "KeywordHits"])

    doc.save(output_path)
    return str(output_path)


def _add_table(doc: Document, df: pd.DataFrame, columns: list[str]) -> None:
    table = doc.add_table(rows=1, cols=len(columns))
    for idx, col in enumerate(columns):
        table.rows[0].cells[idx].text = col
    if df.empty:
        cells = table.add_row().cells
        cells[0].text = "无"
        return
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for idx, col in enumerate(columns):
            value = row.get(col, "")
            if isinstance(value, float):
                cells[idx].text = f"{value:.6g}"
            else:
                cells[idx].text = str(value)
