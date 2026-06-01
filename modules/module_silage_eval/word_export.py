from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document


def export_word_report(ranking: pd.DataFrame, weights: pd.DataFrame, directions: dict[str, str], output_path: str | Path) -> str:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading("青贮综合评价报告", level=1)
    doc.add_paragraph("本报告汇总指标方向、熵权结果和综合排名。")

    doc.add_heading("指标方向表", level=2)
    direction_table = doc.add_table(rows=1, cols=2)
    direction_table.rows[0].cells[0].text = "指标"
    direction_table.rows[0].cells[1].text = "方向"
    for name, direction in directions.items():
        cells = direction_table.add_row().cells
        cells[0].text = str(name)
        cells[1].text = "正向" if direction == "positive" else "负向"

    doc.add_heading("熵权表", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "指标"
    table.rows[0].cells[1].text = "权重"
    for _, row in weights.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(row["Indicator"])
        cells[1].text = f"{float(row['Weight']):.6f}"

    doc.add_heading("综合排名表", level=2)
    rank_table = doc.add_table(rows=1, cols=3)
    headers = ["处理", "综合得分", "排名"]
    for idx, header in enumerate(headers):
        rank_table.rows[0].cells[idx].text = header
    for _, row in ranking.iterrows():
        cells = rank_table.add_row().cells
        cells[0].text = str(row["Treatment"])
        cells[1].text = f"{float(row['Score']):.6f}"
        cells[2].text = str(int(row["Rank"]))

    doc.save(output_path)
    return str(output_path)
