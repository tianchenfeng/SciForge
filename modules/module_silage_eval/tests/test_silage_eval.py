from __future__ import annotations

from pathlib import Path
import sys

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from modules.module_silage_eval.service import SilageEvalService


def main():
    base = Path(__file__).resolve().parent
    input_file = base / "demo_silage_eval.xlsx"
    output_dir = base / "output_silage_eval"
    df = pd.DataFrame({
        "Treatment": ["T1", "T1", "T1", "T2", "T2", "T2", "T3", "T3", "T3"],
        "pH": [4.2, 4.1, 4.3, 3.8, 3.9, 3.7, 4.0, 4.1, 3.9],
        "LA": [55, 56, 54, 70, 72, 68, 62, 63, 61],
        "NH3-N": [8.2, 8.0, 8.4, 6.1, 6.0, 6.2, 7.5, 7.3, 7.7],
        "NDF": [48, 49, 47, 45, 44, 46, 46, 47, 45],
        "ADF": [31, 32, 30, 29, 28, 30, 30, 31, 29],
        "NDFD": [50, 51, 49, 58, 57, 59, 55, 54, 56],
        "ADFD": [40, 41, 39, 48, 47, 49, 45, 44, 46],
        "ConstantMetric": [1, 1, 1, 1, 1, 1, 1, 1, 1],
        "Remark": ["a", "b", "c", "d", "e", "f", "g", "h", "i"],
        "EmptyColumn": [None] * 9,
    })
    df.to_excel(input_file, index=False)
    result = SilageEvalService.run(str(input_file), str(output_dir), logger=print)
    assert Path(result["excel"]).exists()
    assert Path(result["word"]).exists()
    assert len(result["ranking"]) == 3
    assert len(result["treatment_mean"]) == 3
    assert result["directions"]["NDFD"] == "positive"
    assert result["directions"]["ADFD"] == "positive"
    assert result["directions"]["NDF"] == "negative"
    assert result["directions"]["ADF"] == "negative"
    assert result["directions"]["NH3-N"] == "negative"

    workbook = load_workbook(result["excel"], data_only=True)
    for sheet_name in ["raw_data", "treatment_mean", "directions", "normalized", "entropy_weights", "ranking"]:
        assert sheet_name in workbook.sheetnames

    normalized = pd.read_excel(result["excel"], sheet_name="normalized")
    assert "Remark" not in normalized.columns
    assert "EmptyColumn" not in normalized.columns
    assert (normalized["ConstantMetric"] == 0.0).all()

    print(f"directions={result['directions']}")
    print(f"treatment_mean_rows={len(result['treatment_mean'])}")
    print(f"ranking_rows={len(result['ranking'])}")
    print("module_silage_eval test passed")


if __name__ == "__main__":
    main()
