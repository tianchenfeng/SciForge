from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _format_p_value(p):
    try:
        p = float(p)
    except Exception:
        return "NA"

    if p != p or p in [float("inf"), float("-inf")]:
        return "NA"

    if p < 0.001:
        return "<0.001"
    return f"{p:.3f}"


def _format_number(value, digits=2):
    try:
        return f"{float(value):.{digits}f}"
    except Exception:
        return "NA"


def _clear_cell(cell):
    cell.text = ""


def _set_cell_text(cell, text, bold=False):
    _clear_cell(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(10.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_value_with_superscript(cell, value_text, letter_text):
    _clear_cell(cell)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run_value = p.add_run(str(value_text))
    run_value.font.size = Pt(10.5)

    if letter_text:
        run_letter = p.add_run(str(letter_text))
        run_letter.font.size = Pt(8)
        run_letter.font.superscript = True

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _extract_value_and_letter(row):
    value_text = _format_number(row.get("mean"), digits=2)
    letter_text = row.get("final_letter", "")
    if letter_text is None:
        letter_text = ""
    return value_text, str(letter_text)


def _anova_p(anova, key):
    if key in anova.index and "PR(>F)" in anova.columns:
        return anova.loc[key, "PR(>F)"]
    return None


def export_word_report(output_path, result, factor1, factor2, value_col):
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{value_col} two-way ANOVA")
    run.font.size = Pt(14)
    run.bold = False

    summary = result["summary"].copy()
    anova = result["anova_table"]

    factor1_levels = list(summary[factor1].drop_duplicates())
    factor2_levels = list(summary[factor2].drop_duplicates())
    sem_value = summary["sem"].mean()

    p_factor1 = _anova_p(anova, f"C({factor1})")
    p_factor2 = _anova_p(anova, f"C({factor2})")
    p_inter = _anova_p(anova, f"C({factor1}):C({factor2})")

    n_cols = 2 + len(factor1_levels) + 4
    table = doc.add_table(rows=1, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    hdr = table.rows[0].cells
    _set_cell_text(hdr[0], "Project", bold=True)
    _set_cell_text(hdr[1], factor2, bold=True)

    for i, lv in enumerate(factor1_levels):
        _set_cell_text(hdr[2 + i], str(lv), bold=True)

    sem_col = 2 + len(factor1_levels)
    _set_cell_text(hdr[sem_col], "SEM", bold=True)
    _set_cell_text(hdr[sem_col + 1], f"P_{factor1}", bold=True)
    _set_cell_text(hdr[sem_col + 2], f"P_{factor2}", bold=True)
    _set_cell_text(hdr[sem_col + 3], f"P_{factor1}\u00d7{factor2}", bold=True)

    for row_idx, b_lv in enumerate(factor2_levels):
        cells = table.add_row().cells
        _set_cell_text(cells[0], value_col if row_idx == 0 else "")
        _set_cell_text(cells[1], str(b_lv))

        for i, a_lv in enumerate(factor1_levels):
            sub = summary[
                (summary[factor1].astype(str) == str(a_lv))
                & (summary[factor2].astype(str) == str(b_lv))
            ]

            if sub.empty:
                _set_cell_text(cells[2 + i], "NA")
                continue

            value_text, letter_text = _extract_value_and_letter(sub.iloc[0])
            _set_value_with_superscript(cells[2 + i], value_text, letter_text)

        if row_idx == 0:
            _set_cell_text(cells[sem_col], _format_number(sem_value, digits=3))
            _set_cell_text(cells[sem_col + 1], _format_p_value(p_factor1))
            _set_cell_text(cells[sem_col + 2], _format_p_value(p_factor2))
            _set_cell_text(cells[sem_col + 3], _format_p_value(p_inter))
        else:
            _set_cell_text(cells[sem_col], "")
            _set_cell_text(cells[sem_col + 1], "")
            _set_cell_text(cells[sem_col + 2], "")
            _set_cell_text(cells[sem_col + 3], "")

    doc.save(output_path)


def export_one_way_word_report(output_path, result, factor, value_col):
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{value_col} one-way ANOVA")
    run.font.size = Pt(14)
    run.bold = False

    summary = result["summary"].copy()
    anova = result["anova_table"]
    factor_levels = list(summary[factor].astype(str))

    p_factor = _anova_p(anova, f"C({factor})")
    sem_value = summary["sem"].mean()

    n_cols = 1 + len(factor_levels) + 2
    table = doc.add_table(rows=1, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    hdr = table.rows[0].cells
    _set_cell_text(hdr[0], "Project", bold=True)

    for i, lv in enumerate(factor_levels):
        _set_cell_text(hdr[i + 1], str(lv), bold=True)

    sem_col = 1 + len(factor_levels)
    _set_cell_text(hdr[sem_col], "SEM", bold=True)
    _set_cell_text(hdr[sem_col + 1], "P", bold=True)

    cells = table.add_row().cells
    _set_cell_text(cells[0], value_col)

    for i, (_, row) in enumerate(summary.iterrows()):
        value_text, letter_text = _extract_value_and_letter(row)
        _set_value_with_superscript(cells[i + 1], value_text, letter_text)

    _set_cell_text(cells[sem_col], _format_number(sem_value, digits=3))
    _set_cell_text(cells[sem_col + 1], _format_p_value(p_factor))

    doc.save(output_path)


def export_one_way_batch_word_report(output_path, batch_result, factor):
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Batch one-way ANOVA")
    run.font.size = Pt(14)
    run.bold = False

    paper_table = batch_result["paper_table"]
    columns = list(paper_table.columns)

    table = doc.add_table(rows=1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    hdr = table.rows[0].cells
    for i, column in enumerate(columns):
        _set_cell_text(hdr[i], column, bold=True)

    for _, paper_row in paper_table.iterrows():
        metric = paper_row["Project"]
        result = batch_result["results"][metric]
        summary = result["summary"].copy()
        cells = table.add_row().cells

        _set_cell_text(cells[0], metric)

        for i, level in enumerate(columns[1:-2], start=1):
            sub = summary[summary[factor].astype(str) == str(level)]
            if sub.empty:
                _set_cell_text(cells[i], "NA")
                continue

            value_text, letter_text = _extract_value_and_letter(sub.iloc[0])
            _set_value_with_superscript(cells[i], value_text, letter_text)

        _set_cell_text(cells[-2], paper_row["SEM"])
        _set_cell_text(cells[-1], paper_row["P"])

    doc.save(output_path)


def export_two_way_batch_word_report(output_path, batch_result, factor1, factor2):
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Batch two-way ANOVA")
    run.font.size = Pt(14)
    run.bold = False

    paper_table = batch_result["paper_table"]
    columns = list(paper_table.columns)

    table = doc.add_table(rows=1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    hdr = table.rows[0].cells
    for i, column in enumerate(columns):
        _set_cell_text(hdr[i], column, bold=True)

    current_metric = None

    for _, paper_row in paper_table.iterrows():
        if paper_row["Project"]:
            current_metric = paper_row["Project"]

        metric = current_metric
        result = batch_result["results"][metric]
        summary = result["summary"].copy()
        cells = table.add_row().cells

        _set_cell_text(cells[0], paper_row["Project"])
        _set_cell_text(cells[1], paper_row[factor2])

        for i, level in enumerate(columns[2:-4], start=2):
            sub = summary[
                (summary[factor1].astype(str) == str(level))
                & (summary[factor2].astype(str) == str(paper_row[factor2]))
            ]
            if sub.empty:
                _set_cell_text(cells[i], "NA")
                continue

            value_text, letter_text = _extract_value_and_letter(sub.iloc[0])
            _set_value_with_superscript(cells[i], value_text, letter_text)

        _set_cell_text(cells[-4], paper_row["SEM"])
        _set_cell_text(cells[-3], paper_row[f"P_{factor1}"])
        _set_cell_text(cells[-2], paper_row[f"P_{factor2}"])
        _set_cell_text(cells[-1], paper_row[f"P_{factor1}\u00d7{factor2}"])

    doc.save(output_path)
